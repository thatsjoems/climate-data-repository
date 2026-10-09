"""Portfolio breakdown (dashboard charts): grouping, Others folding, drill-down, tenant isolation."""
import pytest

from app.models.models import RoleEnum, SubmissionRecord, SubmissionStatus
from tests.conftest import auth_header, login, make_institution, make_user
from tests.test_rbac_and_isolation import _seed_submission_for


def _get(client, username, path):
    token = login(client, username).json()["access_token"]
    return client.get(path, headers=auth_header(token))


def _setup(db):
    a = make_institution(db, code="BANK-A", name="Bank A Ltd")
    b = make_institution(db, code="BANK-B", name="Bank B Ltd")
    _seed_submission_for(db, a, region="Dodoma", district="Chamwino", amount=1_000_000.0)
    _seed_submission_for(db, a, region="Mwanza", district="Nyamagana", amount=3_000_000.0)
    _seed_submission_for(db, b, region="Mwanza", district="Ilemela", amount=1_000_000.0)
    # A submission still awaiting review must never appear in any chart. It belongs to a THIRD institution on purpose:
    # _seed_submission_for adds a second call for the same institution to its existing submission and ignores `status`.
    c = make_institution(db, code="BANK-C", name="Bank C Ltd")
    _seed_submission_for(db, c, region="Arusha", amount=9_000_000.0, status=SubmissionStatus.VALID)
    make_user(db, role=RoleEnum.BOT_USER, username="bot1")
    make_user(db, role=RoleEnum.INSTITUTION_USER, institution=a, username="usera")
    return a, b


def test_breakdown_by_region_is_sorted_and_sums_to_100(client, db_session):
    _setup(db_session)
    data = _get(client, "bot1", "/api/analytics/portfolio-breakdown?group_by=region").json()
    assert [(d["label"], d["value"]) for d in data] == [("Mwanza", 4_000_000.0), ("Dodoma", 1_000_000.0)]
    assert sum(d["share_pct"] for d in data) == pytest.approx(100.0)
    assert "Arusha" not in {d["label"] for d in data}  # not approved
    by_bank = _get(client, "bot1", "/api/analytics/portfolio-breakdown?group_by=institution").json()
    assert "Bank C Ltd" not in {d["label"] for d in by_bank}


def test_breakdown_by_institution(client, db_session):
    _setup(db_session)
    data = _get(client, "bot1", "/api/analytics/portfolio-breakdown?group_by=institution").json()
    assert {d["label"]: d["value"] for d in data} == {"Bank A Ltd": 4_000_000.0, "Bank B Ltd": 1_000_000.0}


def test_drill_down_region_to_district(client, db_session):
    _setup(db_session)
    data = _get(client, "bot1", "/api/analytics/portfolio-breakdown?group_by=district&filter_region=Mwanza").json()
    assert {d["label"] for d in data} == {"Nyamagana", "Ilemela"}
    one = _get(client, "bot1", "/api/analytics/portfolio-breakdown?group_by=district&filter_region=Mwanza&filter_district=Ilemela").json()
    assert [(d["label"], d["value"]) for d in one] == [("Ilemela", 1_000_000.0)]


def test_small_groups_fold_into_others(client, db_session):
    _setup(db_session)
    data = _get(client, "bot1", "/api/analytics/portfolio-breakdown?group_by=district&limit=2").json()
    assert len(data) == 2 and data[-1]["label"] == "Others"
    assert sum(d["value"] for d in data) == pytest.approx(5_000_000.0)


def test_records_metric_counts_rows(client, db_session):
    _setup(db_session)
    data = _get(client, "bot1", "/api/analytics/portfolio-breakdown?group_by=region&metric=records").json()
    assert {d["label"]: d["value"] for d in data} == {"Mwanza": 2.0, "Dodoma": 1.0}


def test_institution_user_sees_only_own_data(client, db_session):
    _setup(db_session)
    data = _get(client, "usera", "/api/analytics/portfolio-breakdown?group_by=institution").json()
    assert [d["label"] for d in data] == ["Bank A Ltd"]
    # An institution filter pointing at another bank cannot widen the scope.
    other = db_session.query(SubmissionRecord).first()  # noqa: F841
    from app.models.models import Institution
    b = db_session.query(Institution).filter_by(code="BANK-B").one()
    data = _get(client, "usera", f"/api/analytics/portfolio-breakdown?group_by=institution&filter_institution_id={b.id}").json()
    assert all(d["label"] != "Bank B Ltd" for d in data)


def test_unknown_dimension_or_metric_is_refused(client, db_session):
    _setup(db_session)
    assert _get(client, "bot1", "/api/analytics/portfolio-breakdown?group_by=password_hash").status_code == 422
    assert _get(client, "bot1", "/api/analytics/portfolio-breakdown?group_by=region&metric=drop").status_code == 422


def test_system_admin_is_refused(client, db_session):
    _setup(db_session)
    make_user(db_session, role=RoleEnum.SYSTEM_ADMIN, username="adm1")
    assert _get(client, "adm1", "/api/analytics/portfolio-breakdown?group_by=region").status_code == 403


def test_excel_report_carries_the_portfolio_breakdown(client, db_session):
    import io
    from openpyxl import load_workbook
    _setup(db_session)
    token = login(client, "bot1").json()["access_token"]
    res = client.get("/api/reports/summary.xlsx", headers=auth_header(token))
    assert res.status_code == 200
    ws = load_workbook(io.BytesIO(res.content))["Portfolio Breakdown"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert ["Loan by bank", "Bank A Ltd", 4_000_000.0, 80.0] in rows


def test_pdf_report_still_builds_with_the_portfolio_section(client, db_session):
    _setup(db_session)
    token = login(client, "bot1").json()["access_token"]
    res = client.get("/api/reports/summary.pdf", headers=auth_header(token))
    assert res.status_code == 200 and res.content.startswith(b"%PDF")
