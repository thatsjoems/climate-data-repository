"""
The load-test tool (scripts/load_test.py), which is meant for the STAGING stack only.

What matters: it refuses to run anywhere that is not staging; its figures are computed correctly; and, above all, the synthetic files it uploads
are ones the REAL upload endpoint accepts as valid. If they were not, every upload figure of a load test would measure the rejection of bad data,
not the system. The last test below proves it against the real application.
"""
import importlib.util
import io
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.core.password_policy import validate_password_strength
from app.services import geo_lookup
from app.services.template_generator import COLUMNS
from tests.conftest import auth_header, login
from tests.test_upload_and_workflow import _setup_institution_user

BACKEND = Path(__file__).resolve().parent.parent
_spec = importlib.util.spec_from_file_location("load_test_tool", BACKEND / "scripts" / "load_test.py")
lt = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(lt)


# ------------------------------------------------------------------ never against production
def test_it_runs_only_where_staging_says_yes(monkeypatch):
    monkeypatch.setenv("CDR_ALLOW_LOAD_TEST", "yes")
    lt.require_staging()                                                     # allowed: no exception


@pytest.mark.parametrize("value", [None, "", "true", "1", "YES", "no"])
def test_it_refuses_everywhere_else(monkeypatch, value):
    if value is None:
        monkeypatch.delenv("CDR_ALLOW_LOAD_TEST", raising=False)
    else:
        monkeypatch.setenv("CDR_ALLOW_LOAD_TEST", value)
    with pytest.raises(SystemExit) as stopped:
        lt.require_staging()
    assert "REFUSED" in str(stopped.value) and "never run against production" in str(stopped.value)


def test_the_missing_setup_message_explains_the_unprivileged_user_when_the_tool_fell_back_to_tmp(tmp_path):
    uploads = tmp_path / "uploads"
    uploads.mkdir()
    assert "-u cdr" in lt.not_set_up_message(Path("/tmp"), uploads)                      # the uploads volume exists but the tool looked in /tmp
    assert "-u cdr" not in lt.not_set_up_message(uploads / "_loadtest", uploads)         # it looked in the right place: nothing to explain
    assert "-u cdr" not in lt.not_set_up_message(Path("/tmp"), tmp_path / "no-such-folder")   # no uploads volume at all (development)
    assert lt.not_set_up_message(uploads / "_loadtest", uploads).startswith("Not set up yet")


def test_the_passwords_of_the_test_users_are_kept_where_they_survive_a_restart(tmp_path):
    assert lt.default_dir(tmp_path) == tmp_path / "_loadtest"                # the uploads volume of staging survives the container being recreated
    assert lt.default_dir(tmp_path / "missing") == Path("/tmp")              # anywhere else: a temporary folder


# ------------------------------------------------------------------ the figures
@pytest.mark.parametrize("p, expected", [(50, 5), (95, 10), (99, 10), (100, 10), (1, 1)])
def test_percentile_is_the_nearest_rank(p, expected):
    assert lt.percentile(list(range(1, 11)), p) == expected


def test_percentile_of_nothing_is_zero_not_an_error():
    assert lt.percentile([], 95) == 0.0


def test_the_summary_reports_milliseconds_errors_and_rate():
    samples = [(True, 0.010), (True, 0.020), (True, 0.030), (False, 0.500)]
    s = lt.summarise(samples, elapsed=2.0)
    assert s["requests"] == 4 and s["errors"] == 1 and s["error_rate"] == 0.25 and s["per_second"] == 2.0
    assert s["min_ms"] == 10 and s["max_ms"] == 500 and s["p50_ms"] == 20 and s["p95_ms"] == 500 and round(s["mean_ms"], 1) == 140.0


def test_an_empty_summary_has_no_division_by_zero():
    s = lt.summarise([], elapsed=0)
    assert s["requests"] == 0 and s["per_second"] == 0.0 and s["error_rate"] == 0.0


def test_the_recorder_keeps_the_most_frequent_errors():
    recorder = lt.Recorder()
    for _ in range(5):
        recorder.add("read", False, 0.1, "HTTP 500 boom")
    recorder.add("read", False, 0.1, "HTTP 502")
    recorder.add("read", True, 0.1)
    summary = recorder.summary("read", 1.0)
    assert summary["errors"] == 6 and summary["error_samples"][0] == ("HTTP 500 boom", 5)


# ------------------------------------------------------------------ memory of the container
def test_the_memory_reader_understands_the_current_and_the_older_layout(tmp_path):
    (tmp_path / "memory.current").write_text("2097152\n")
    (tmp_path / "memory.peak").write_text("4194304\n")
    (tmp_path / "memory.max").write_text("max\n")                                   # "max" means no limit was set
    assert lt.container_memory(tmp_path) == {"current": 2097152, "peak": 4194304, "limit": None}
    old = tmp_path / "old"
    (old / "memory").mkdir(parents=True)
    (old / "memory" / "memory.usage_in_bytes").write_text("1000")
    (old / "memory" / "memory.max_usage_in_bytes").write_text("3000")
    (old / "memory" / "memory.limit_in_bytes").write_text("2000000")
    assert lt.container_memory(old) == {"current": 1000, "peak": 3000, "limit": 2000000}


def test_the_memory_reader_gives_nothing_rather_than_an_error_when_the_system_does_not_say(tmp_path):
    assert lt.container_memory(tmp_path / "nowhere") == {"current": None, "peak": None, "limit": None}
    (tmp_path / "memory.current").write_text("not a number")
    assert lt.container_memory(tmp_path)["current"] is None


def test_an_unset_limit_written_as_a_huge_number_is_no_limit(tmp_path):
    (tmp_path / "memory.limit_in_bytes").write_text("9223372036854771712")
    (tmp_path / "memory").mkdir()
    (tmp_path / "memory" / "memory.limit_in_bytes").write_text("9223372036854771712")
    assert lt.container_memory(tmp_path)["limit"] is None


# ------------------------------------------------------------------ the synthetic data
def test_generated_passwords_meet_the_policy_and_are_not_repeated():
    passwords = {lt.strong_password() for _ in range(100)}
    assert len(passwords) == 100 and all(validate_password_strength(p) == [] for p in passwords)


def test_generated_loans_are_unique_real_and_within_the_rules_of_the_database():
    rows = lt.make_rows(400, "T1", seed=7)
    assert len(rows) == 400 and len({r["loan_id"] for r in rows}) == 400 and len({r["customer_id"] for r in rows}) == 400
    for r in rows:
        assert geo_lookup.is_valid_district(r["region"], r["district"])
        assert 0 < r["outstanding_principal_tzs"] <= r["loan_amount_tzs"]
        assert 0 <= r["annual_interest_rate"] <= 100 and r["collateral_value_tzs"] > 0


def test_the_same_seed_gives_the_same_loans_and_another_seed_different_ones():
    assert lt.make_rows(50, "T", seed=1) == lt.make_rows(50, "T", seed=1)
    assert lt.make_rows(50, "T", seed=1) != lt.make_rows(50, "T", seed=2)


def test_the_workbook_has_the_official_header_and_every_row():
    rows = lt.make_rows(120, "T2", seed=3)
    sheet = load_workbook(io.BytesIO(lt.build_workbook(rows)))[lt.SHEET]
    data = list(sheet.iter_rows(values_only=True))
    assert list(data[0]) == [label for _, label in COLUMNS] and len(data) == 121
    fields = [f for f, _ in COLUMNS]
    assert [r[fields.index("loan_id")] for r in data[1:]] == [r["loan_id"] for r in rows]


def test_a_multipart_body_carries_the_fields_and_the_file():
    body, content_type = lt.encode_multipart({"reporting_period": "2026-Q1"}, {"file": ("a.xlsx", b"PK-data", lt.XLSX)})
    boundary = content_type.split("boundary=")[1].encode()
    assert body.count(b"--" + boundary) == 3 and body.rstrip().endswith(b"--" + boundary + b"--")
    assert b'name="reporting_period"\r\n\r\n2026-Q1' in body and b'filename="a.xlsx"' in body and b"PK-data" in body


# ------------------------------------------------------------------ the one that matters: the real endpoint accepts the generated file
def test_the_real_upload_accepts_a_generated_file_as_valid(client, db_session):
    _setup_institution_user(db_session)
    token = login(client, "inst_user").json()["access_token"]
    workbook = lt.build_workbook(lt.make_rows(150, "REAL", seed=11))
    res = client.post("/api/submissions/upload", data={"reporting_period": "2026-Q1"},
                      files={"file": ("load_1.xlsx", workbook, lt.XLSX)}, headers=auth_header(token))
    assert res.status_code == 201, res.text
    body = res.json()
    assert body["status"] == "VALID" and body["valid_records"] == 150 and body.get("invalid_records", 0) == 0
