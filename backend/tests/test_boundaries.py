"""
Coordinates are checked against the official ward boundaries (NBS, 2022 census): a loan or collateral whose
latitude/longitude is not in the region or district the row names is rejected; a different ward is a warning.

The first half tests the boundary service on its own (no database); the second half tests what an upload does with it.
"""
import json
from pathlib import Path

import pytest

from app.models.models import SubmissionRecord
from app.services import boundary_service as bs
from app.services.validation_service import validate_excel_file
from tests.conftest import login
from tests.test_upload_and_workflow import VALID_ROW, _setup_institution_user, _upload, build_test_excel

BEREKO = (-4.4566, 35.7547)         # inside Bereko ward, Kondoa district, Dodoma (the template's example row)
DODOMA_CITY = (-6.163, 35.7516)     # Dodoma region, Dodoma district
MPANDA = (-6.3677, 31.2626)         # Katavi region, Mpanda district
ILALA = (-6.8235, 39.2695)          # Dar es Salaam, Ilala
SEA = (-6.0, 40.9)                  # the Indian Ocean, east of Zanzibar


# --------------------------------------------------------------------------------------------- the boundary data
def test_the_boundary_data_holds_every_ward_region_and_district_of_the_census():
    b = bs._load()
    assert bs.is_available()
    assert bs.ward_count() == 4344
    assert len(b.by_region) == 31
    assert len(b.by_district) == 150


def test_every_region_and_district_in_the_repository_geography_exists_in_the_boundaries():
    geography = json.loads((Path(bs.__file__).resolve().parent.parent / "data" / "tanzania_geography.json").read_text(encoding="utf-8"))
    for region, districts in geography.items():
        for district in districts:
            level, wards = bs._claimed_wards_cached(region.strip(), district.strip(), "")
            assert level == "district" and wards, f"{region} / {district} has no boundary"


def test_almost_every_ward_in_the_repository_geography_is_found_in_the_boundaries():
    """The two lists come from different sources. A ward that is not found is not rejected: it is checked at district level."""
    geography = json.loads((Path(bs.__file__).resolve().parent.parent / "data" / "tanzania_geography.json").read_text(encoding="utf-8"))
    total = found = 0
    for region, districts in geography.items():
        for district, wards in districts.items():
            for ward in wards:
                total += 1
                found += bs._claimed_wards_cached(region.strip(), district.strip(), ward.strip())[0] == "ward"
    assert total == 4342
    assert found / total > 0.995, f"only {found} of {total} wards matched"


# --------------------------------------------------------------------------------------------- where is a point
def test_a_point_is_found_in_its_own_ward():
    assert bs.locate(*BEREKO) == bs.Place("Dodoma", "Kondoa", "Bereko")
    assert bs.locate(*MPANDA).region == "Katavi"
    assert bs.locate(*ILALA).district == "Ilala"


def test_a_point_in_the_sea_is_in_no_ward():
    assert bs.locate(*SEA) is None


def test_a_point_inside_the_named_place_is_accepted_at_each_level():
    assert bs.check_location("Dodoma", "Kondoa", "Bereko", *BEREKO).kind == "inside"
    assert bs.check_location("Dodoma", "Kondoa", "", *BEREKO).kind == "inside"
    assert bs.check_location("Dodoma", "", "", *BEREKO).kind == "inside"


def test_katavi_coordinates_under_morogoro_are_a_region_mismatch():
    finding = bs.check_location("Morogoro", "Morogoro", "", *MPANDA)
    assert finding.kind == "region"
    assert finding.actual.region == "Katavi" and finding.actual.district == "Mpanda"


def test_the_right_region_but_the_wrong_district_is_a_district_mismatch():
    finding = bs.check_location("Dodoma", "Kondoa", "", *DODOMA_CITY)
    assert finding.kind == "district" and finding.actual.district == "Dodoma"


def test_the_right_district_but_the_wrong_ward_is_a_ward_mismatch():
    finding = bs.check_location("Dodoma", "Kondoa", "Bereko", -4.8, 35.8)
    assert finding.kind == "ward" and finding.actual.district == "Kondoa" and finding.actual.ward != "Bereko"


def test_a_point_in_the_sea_is_outside_not_a_mismatch():
    assert bs.check_location("Dar es Salaam", "Ilala", "", *SEA).kind == "outside"


def test_spelling_case_and_spacing_do_not_matter():
    assert bs.check_location("dodoma", " KONDOA ", "bereko", *BEREKO).kind == "inside"
    assert bs.check_location("Dodoma", "Kondoa", "  Bereko", *BEREKO).kind == "inside"


def test_a_place_the_boundaries_do_not_know_cannot_be_judged():
    assert bs.check_location("Narnia", "Nowhere", "", *BEREKO) is None


def test_an_unmatched_ward_is_checked_as_its_district():
    # A ward name that the boundary file does not hold (here one that does not exist) is judged by its district instead.
    assert bs._claimed_wards_cached("Dodoma", "Kondoa", "Not A Real Ward")[0] == "district"


def test_a_point_just_over_a_boundary_is_accepted_but_one_far_over_it_is_not():
    """500 m of tolerance: simplified shapes, different survey datums and a phone's GPS error never refuse a farm on the edge."""
    b = bs._load()
    ward = b.by_ward[("dodoma", "kondoa", "bereko")][0]
    lat, lon = BEREKO
    step = 0.0005                                    # about 55 m
    while ward.contains(lon * b.scale, lat * b.scale):
        lat += step                                  # walk north until Bereko ends
    assert bs.check_location("Dodoma", "Kondoa", "Bereko", lat + 0.002, lon).kind == "inside"      # ~220 m over the edge
    far = bs.check_location("Dodoma", "Kondoa", "Bereko", lat + 0.05, lon)                         # ~5.5 km over it
    assert far.kind != "inside"


# --------------------------------------------------------------------------------------------- typing mistakes
def test_a_forgotten_minus_sign_is_named_in_the_hint():
    hint = bs.hint_for_typing_mistake("Dodoma", "Kondoa", "", 4.4566, 35.7547)
    assert hint and "minus sign" in hint


def test_swapped_latitude_and_longitude_are_named_in_the_hint():
    hint = bs.hint_for_typing_mistake("Dodoma", "Kondoa", "", 35.7547, -4.4566)
    assert hint and "swapped" in hint


def test_there_is_no_hint_when_no_simple_mistake_explains_it():
    assert bs.hint_for_typing_mistake("Morogoro", "Morogoro", "", *MPANDA) is None


# --------------------------------------------------------------------------------------------- without the data
def test_without_the_boundary_file_nothing_is_checked_and_nothing_breaks(monkeypatch, tmp_path):
    monkeypatch.setattr(bs, "_DATA_PATH", tmp_path / "missing.json.gz")
    bs._load.cache_clear()
    bs._claimed_wards_cached.cache_clear()
    try:
        assert bs.is_available() is False
        assert bs.check_location("Morogoro", "Morogoro", "", *MPANDA) is None
        assert bs.locate(*MPANDA) is None
    finally:
        monkeypatch.undo()
        bs._load.cache_clear()
        bs._claimed_wards_cached.cache_clear()


# --------------------------------------------------------------------------------------------- what an upload does
def _validate(row):
    records, issues = validate_excel_file(build_test_excel([dict(VALID_ROW, **row)]), "data.xlsx")
    return records, issues


def _where(row):
    return dict(region="Dodoma", district="Kondoa", ward="Bereko", loan_latitude=BEREKO[0], loan_longitude=BEREKO[1], **row)


def test_a_loan_in_the_place_it_names_is_valid_and_has_no_coordinate_issue():
    records, issues = _validate(_where({}))
    assert [i for i in issues if i.column_name == "loan_latitude"] == []
    assert records[0]["is_valid"] is True


def test_a_loan_in_another_region_is_rejected_with_the_place_it_is_really_in():
    records, issues = _validate(dict(region="Morogoro", district="Morogoro", loan_latitude=MPANDA[0], loan_longitude=MPANDA[1]))
    problem = [i for i in issues if i.column_name == "loan_latitude"]
    assert len(problem) == 1 and problem[0].severity == "ERROR"
    assert "Katavi region" in problem[0].description and "Mpanda district" in problem[0].description
    assert records[0]["is_valid"] is False


def test_a_loan_in_another_district_of_the_same_region_is_rejected():
    records, issues = _validate(dict(region="Dodoma", district="Kondoa", loan_latitude=DODOMA_CITY[0], loan_longitude=DODOMA_CITY[1]))
    problem = [i for i in issues if i.column_name == "loan_latitude"]
    assert problem and problem[0].severity == "ERROR"
    assert records[0]["is_valid"] is False


def test_a_loan_in_another_ward_of_the_same_district_is_only_a_warning():
    records, issues = _validate(dict(region="Dodoma", district="Kondoa", ward="Bereko", loan_latitude=-4.8, loan_longitude=35.8))
    problem = [i for i in issues if i.column_name == "loan_latitude"]
    assert problem and problem[0].severity == "WARNING"
    assert records[0]["is_valid"] is True


def test_the_message_names_a_forgotten_minus_sign():
    records, issues = _validate(dict(region="Dodoma", district="Kondoa", loan_latitude=4.4566, loan_longitude=35.7547))
    problem = [i for i in issues if i.column_name == "loan_latitude"]
    assert problem and "minus sign" in problem[0].description
    assert problem[0].severity == "WARNING"            # the point is in no ward at all (South Sudan): a warning, not a verdict
    assert records[0]["is_valid"] is True


def test_a_point_in_the_sea_is_a_warning_not_a_rejection():
    records, issues = _validate(dict(region="Dar es Salaam", district="Ilala", loan_latitude=SEA[0], loan_longitude=SEA[1]))
    problem = [i for i in issues if i.column_name == "loan_latitude"]
    assert problem and all(i.severity == "WARNING" for i in problem)
    assert records[0]["is_valid"] is True


def test_the_collateral_location_is_checked_in_the_same_way():
    records, issues = _validate(dict(
        region="Dodoma", district="Kondoa", ward="Bereko", loan_latitude=BEREKO[0], loan_longitude=BEREKO[1],
        collateral_region="Morogoro", collateral_district="Morogoro",
        collateral_latitude=MPANDA[0], collateral_longitude=MPANDA[1],
    ))
    problem = [i for i in issues if i.column_name == "collateral_latitude"]
    assert problem and problem[0].severity == "ERROR" and "Collateral location" in problem[0].description
    assert records[0]["is_valid"] is False


def test_a_row_without_coordinates_is_not_checked():
    records, issues = _validate(dict(region="Dodoma", district="Kondoa"))
    assert [i for i in issues if i.column_name in ("loan_latitude", "loan_longitude")] == []


def test_the_templates_own_example_row_passes_the_check():
    from app.services.template_generator import COLUMNS  # noqa: F401  (the example is written into the template)
    import io
    from openpyxl import load_workbook
    from app.services.template_generator import generate_loan_collateral_template as build_template
    ws = load_workbook(io.BytesIO(build_template()))["Loan_Collateral_Data"]
    labels = {c.value: c.column for c in ws[9] if c.value}
    row = 10
    lat = ws.cell(row=row, column=next(col for label, col in labels.items() if "Latitude" in str(label) and "Collateral" not in str(label))).value
    lon = ws.cell(row=row, column=next(col for label, col in labels.items() if "Longitude" in str(label) and "Collateral" not in str(label))).value
    assert bs.check_location("Dodoma", "Kondoa", "Bereko", float(lat), float(lon)).kind == "inside"


def test_an_upload_with_a_mismatched_loan_is_stored_invalid_and_the_good_one_is_not_affected(client, db_session):
    _setup_institution_user(db_session)
    token = login(client, "inst_user").json()["access_token"]
    rows = [
        dict(VALID_ROW, loan_id="LN-OK", region="Dodoma", district="Kondoa", ward="Bereko",
             loan_latitude=BEREKO[0], loan_longitude=BEREKO[1]),
        dict(VALID_ROW, loan_id="LN-BAD", region="Morogoro", district="Morogoro",
             loan_latitude=MPANDA[0], loan_longitude=MPANDA[1]),
    ]
    res = _upload(client, token, rows)
    assert res.status_code == 201, res.text
    body = res.json()
    assert body["status"] == "INVALID" and body["invalid_records"] == 1
    good = db_session.query(SubmissionRecord).filter(SubmissionRecord.loan_id == "LN-OK").one()
    bad = db_session.query(SubmissionRecord).filter(SubmissionRecord.loan_id == "LN-BAD").one()
    assert good.is_valid is True and bad.is_valid is False
