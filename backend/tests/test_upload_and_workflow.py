"""
Upload/validation and workflow tests.

Covers acceptance criteria: invalid submissions must never be approvable,
reviewers can't approve their own upload, and a corrected resubmission must
supersede the earlier attempt so analytics never double-count it.
"""
import io
from openpyxl import Workbook

from app.models.models import RoleEnum, SubmissionStatus
from tests.conftest import make_institution, make_user, login, auth_header

REQUIRED_HEADERS = [
    "loan_id", "borrower_name", "loan_amount_tzs", "collateral_type",
    "collateral_value_tzs", "region", "district", "reporting_period",
    "climate_hazard_exposure",
]


def build_test_excel(rows: list[dict], reporting_period: str | None = None, sheet_name="Loan_Collateral_Data") -> bytes:
    """Builds a minimal, valid .xlsx in memory matching the CDR template structure."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    for col_idx, header in enumerate(REQUIRED_HEADERS, start=1):
        ws.cell(row=1, column=col_idx, value=header)
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, header in enumerate(REQUIRED_HEADERS, start=1):
            value = row.get(header, "")
            if header == "reporting_period" and reporting_period and header not in row:
                value = reporting_period
            ws.cell(row=row_idx, column=col_idx, value=value)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()


def _upload(client, token, rows, reporting_period="2026-Q1", filename="data.xlsx"):
    file_bytes = build_test_excel(rows, reporting_period=reporting_period)
    return client.post(
        "/api/submissions/upload",
        data={"reporting_period": reporting_period},
        files={"file": (filename, file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=auth_header(token),
    )


VALID_ROW = {
    "loan_id": "LN-1", "borrower_name": "Acme Ltd", "loan_amount_tzs": 5_000_000,
    "collateral_type": "Land Title", "collateral_value_tzs": 8_000_000,
    "region": "Dodoma", "district": "Chamwino District", "reporting_period": "2026-Q1",
    "climate_hazard_exposure": "None",
}


def _setup_institution_user(db_session, username="inst_user", institution_code="BANK-A"):
    inst = make_institution(db_session, code=institution_code, name=f"{institution_code} Ltd")
    make_user(db_session, role=RoleEnum.INSTITUTION_USER, institution=inst, username=username)
    return inst


def test_upload_rejects_non_excel_file(client, db_session):
    _setup_institution_user(db_session)
    token = login(client, "inst_user").json()["access_token"]
    res = client.post(
        "/api/submissions/upload",
        data={"reporting_period": "2026-Q1"},
        files={"file": ("data.txt", b"not an excel file", "text/plain")},
        headers=auth_header(token),
    )
    assert res.status_code == 400


def test_upload_missing_required_columns_marks_invalid(client, db_session):
    _setup_institution_user(db_session)
    token = login(client, "inst_user").json()["access_token"]
    # Build a file whose header row doesn't match the required template columns at all
    wb = Workbook()
    ws = wb.active
    ws.title = "Loan_Collateral_Data"
    ws.cell(row=1, column=1, value="not_a_real_column")
    buf = io.BytesIO()
    wb.save(buf)
    res = client.post(
        "/api/submissions/upload",
        data={"reporting_period": "2026-Q1"},
        files={"file": ("bad.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=auth_header(token),
    )
    assert res.status_code == 201
    assert res.json()["status"] == "INVALID"


def test_upload_invalid_region_and_district_rejected(client, db_session):
    _setup_institution_user(db_session)
    token = login(client, "inst_user").json()["access_token"]
    bad_row = dict(VALID_ROW, region="Narnia", district="Nowhere District")
    res = _upload(client, token, [bad_row])
    assert res.status_code == 201
    assert res.json()["status"] == "INVALID"
    assert res.json()["invalid_records"] == 1


def test_upload_district_must_belong_to_selected_region(client, db_session):
    """A real region with a district that belongs to a DIFFERENT region must fail."""
    _setup_institution_user(db_session)
    token = login(client, "inst_user").json()["access_token"]
    # "Ilala" is a Dar es Salaam district, not a Dodoma one
    bad_row = dict(VALID_ROW, region="Dodoma", district="Ilala")
    res = _upload(client, token, [bad_row])
    assert res.json()["status"] == "INVALID"


def test_reporting_period_mismatch_between_form_and_file_rejected(client, db_session):
    _setup_institution_user(db_session)
    token = login(client, "inst_user").json()["access_token"]
    row = dict(VALID_ROW, reporting_period="2025-Q4")  # file says Q4 2025
    res = _upload(client, token, [row], reporting_period="2026-Q1")  # form says Q1 2026
    assert res.status_code == 201
    assert res.json()["status"] == "INVALID"


def test_duplicate_loan_id_within_same_file_flagged(client, db_session):
    _setup_institution_user(db_session)
    token = login(client, "inst_user").json()["access_token"]
    row2 = dict(VALID_ROW, loan_id="LN-1")  # same loan_id twice
    res = _upload(client, token, [VALID_ROW, row2])
    assert res.json()["status"] == "INVALID"
    assert res.json()["invalid_records"] >= 1


def test_valid_upload_is_accepted(client, db_session):
    _setup_institution_user(db_session)
    token = login(client, "inst_user").json()["access_token"]
    res = _upload(client, token, [VALID_ROW])
    assert res.status_code == 201
    assert res.json()["status"] == "VALID"
    assert res.json()["valid_records"] == 1


def test_invalid_submission_cannot_be_approved(client, db_session):
    inst = _setup_institution_user(db_session)
    inst_token = login(client, "inst_user").json()["access_token"]
    bad_row = dict(VALID_ROW, region="Narnia")
    submission_id = _upload(client, inst_token, [bad_row]).json()["id"]

    make_user(db_session, role=RoleEnum.BOT_USER, username="reviewer1")
    reviewer_token = login(client, "reviewer1").json()["access_token"]

    res = client.post(
        f"/api/submissions/{submission_id}/review",
        json={"decision": "APPROVE"},
        headers=auth_header(reviewer_token),
    )
    assert res.status_code == 400


def test_reviewer_cannot_approve_own_upload(client, db_session):
    """
    Maker-checker defense-in-depth: even though normal role separation (only
    INSTITUTION_USER uploads, only BOT_USER reviews) already makes "the same
    account did both" structurally impossible via the API, the server-side
    guard is tested directly here by constructing that exact data state.
    """
    inst = _setup_institution_user(db_session)
    inst_token = login(client, "inst_user").json()["access_token"]
    submission_id = _upload(client, inst_token, [VALID_ROW]).json()["id"]

    reviewer = make_user(db_session, role=RoleEnum.BOT_USER, username="reviewer1")
    reviewer_token = login(client, "reviewer1").json()["access_token"]

    # Force the data into the "this reviewer is also the submitter" state directly,
    # since that can no longer happen through the API's own role rules.
    from app.models.models import Submission
    submission = db_session.query(Submission).filter(Submission.id == submission_id).first()
    submission.submitted_by_user_id = reviewer.id
    db_session.commit()

    res = client.post(
        f"/api/submissions/{submission_id}/review",
        json={"decision": "APPROVE"},
        headers=auth_header(reviewer_token),
    )
    assert res.status_code == 403


def test_valid_submission_can_be_approved_by_a_different_reviewer(client, db_session):
    _setup_institution_user(db_session)
    inst_token = login(client, "inst_user").json()["access_token"]
    submission_id = _upload(client, inst_token, [VALID_ROW]).json()["id"]

    make_user(db_session, role=RoleEnum.BOT_USER, username="reviewer1")
    reviewer_token = login(client, "reviewer1").json()["access_token"]

    res = client.post(
        f"/api/submissions/{submission_id}/review",
        json={"decision": "APPROVE"},
        headers=auth_header(reviewer_token),
    )
    assert res.status_code == 200
    assert res.json()["status"] == "APPROVED"


def test_already_reviewed_submission_cannot_be_reviewed_again(client, db_session):
    _setup_institution_user(db_session)
    inst_token = login(client, "inst_user").json()["access_token"]
    submission_id = _upload(client, inst_token, [VALID_ROW]).json()["id"]

    make_user(db_session, role=RoleEnum.BOT_USER, username="reviewer1")
    reviewer_token = login(client, "reviewer1").json()["access_token"]
    client.post(f"/api/submissions/{submission_id}/review", json={"decision": "APPROVE"}, headers=auth_header(reviewer_token))

    res = client.post(
        f"/api/submissions/{submission_id}/review",
        json={"decision": "REJECT"},
        headers=auth_header(reviewer_token),
    )
    assert res.status_code == 400


def test_new_upload_supersedes_previous_submission_same_period(client, db_session):
    """
    Acceptance criterion: the same logical exposure must never be double-counted
    because of a corrected/resubmitted upload for the same institution+period.
    """
    _setup_institution_user(db_session)
    token = login(client, "inst_user").json()["access_token"]

    first = _upload(client, token, [dict(VALID_ROW, loan_id="LN-1")], reporting_period="2026-Q1")
    assert first.json()["status"] == "VALID"

    second = _upload(client, token, [dict(VALID_ROW, loan_id="LN-2")], reporting_period="2026-Q1")
    assert second.status_code == 201

    # Re-fetch the first submission - it must now show as SUPERSEDED
    check = client.get(f"/api/submissions/{first.json()['id']}", headers=auth_header(token))
    assert check.json()["status"] == "SUPERSEDED"

    # KPI totals must reflect ONLY the latest (second) submission's exposure, not both
    kpi = client.get("/api/analytics/kpi-summary", headers=auth_header(token))
    assert kpi.json()["total_loan_exposure_tzs"] == VALID_ROW["loan_amount_tzs"]
    assert kpi.json()["total_submissions"] == 2  # both rows exist...
    assert kpi.json()["valid_submissions"] == 1  # ...but only one counts as currently valid/active


def test_system_admin_cannot_upload_submissions(client, db_session):
    """Submissions are exclusively an INSTITUTION_USER action now - not even SYSTEM_ADMIN may upload."""
    inst = make_institution(db_session)
    make_user(db_session, role=RoleEnum.SYSTEM_ADMIN, institution=inst, username="admin1")
    token = login(client, "admin1").json()["access_token"]
    res = _upload(client, token, [VALID_ROW])
    assert res.status_code == 403


def test_institution_can_download_their_own_submission_file(client, db_session):
    _setup_institution_user(db_session)
    token = login(client, "inst_user").json()["access_token"]
    submission_id = _upload(client, token, [VALID_ROW]).json()["id"]

    res = client.get(f"/api/submissions/{submission_id}/download", headers=auth_header(token))
    assert res.status_code == 200


def test_institution_cannot_download_another_institutions_submission_file(client, db_session):
    _setup_institution_user(db_session, username="inst_a_user", institution_code="BANK-A")
    token_a = login(client, "inst_a_user").json()["access_token"]
    submission_id = _upload(client, token_a, [VALID_ROW]).json()["id"]

    _setup_institution_user(db_session, username="inst_b_user", institution_code="BANK-B")
    token_b = login(client, "inst_b_user").json()["access_token"]

    res = client.get(f"/api/submissions/{submission_id}/download", headers=auth_header(token_b))
    assert res.status_code == 403


def test_submission_detail_includes_submitted_records_for_institution_review(client, db_session):
    """Institutions can review the actual data they submitted, not just the error list."""
    _setup_institution_user(db_session)
    token = login(client, "inst_user").json()["access_token"]
    submission_id = _upload(client, token, [VALID_ROW]).json()["id"]

    res = client.get(f"/api/submissions/{submission_id}", headers=auth_header(token))
    assert res.status_code == 200
    records = res.json()["records"]
    assert len(records) == 1
    assert records[0]["loan_id"] == VALID_ROW["loan_id"]
    assert records[0]["is_valid"] is True
